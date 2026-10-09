"""Reading the Brocks et al. (2023) workbooks into one row per sample record and fossil class."""

from __future__ import annotations

import csv
import json
import math
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping

import yaml
from openpyxl import load_workbook

# Columns of Table S5, in order.
OBSERVATION_COLUMNS = (
    "workbook_id",
    "sheet_name",
    "sample_key",
    "source_row",
    "source_sample_id_text",
    "source_sample_id_coordinate",
    "canonical_sample_key",
    "record_kind",
    "included_sample",
    "sample_record",
    "source_rows",
    "cross_workbook_join_applied",
    "geological_era_text",
    "formation_or_basin_text",
    "age_source_coordinate",
    "age_raw_text",
    "age_kind",
    "age_lower_ma",
    "age_upper_ma",
    "age_central_ma",
    "age_operator",
    "age_parse_status",
    "analyte_id",
    "source_coordinate",
    "measurement_channel",
    "denominator_group",
    "class_used",
    "value_used",
    "exclusion_reason",
    "reported_abundance",
    "value_type",
    "qualifier_text",
    "low_signal",
    "cell_value",
)


@dataclass(frozen=True)
class FossilClassColumn:
    column: str
    analyte_id: str
    measurement_channel: str
    denominator_group: str
    used: bool


@dataclass(frozen=True)
class WorkbookLayout:
    workbook_id: str
    file_name: str
    sheet_name: str
    data_start_row: int
    footer_row: int
    sample_id_column: str
    age_column: str
    era_labels: frozenset[str]
    fossil_classes: tuple[FossilClassColumn, ...]


@dataclass(frozen=True)
class Layout:
    workbooks: tuple[WorkbookLayout, ...]
    standard_labels: frozenset[str]
    compilation_labels: frozenset[str]


@dataclass(frozen=True)
class AgeConstraint:
    raw_text: str
    kind: str
    lower_ma: float | None
    upper_ma: float | None
    central_ma: float | None
    operator: str | None
    status: str


def load_layout(path: Path) -> Layout:
    """Read the rows, columns and record labels of the two workbooks."""
    value = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
    workbooks = []
    for book in value["workbooks"]:
        workbooks.append(
            WorkbookLayout(
                workbook_id=str(book["workbook_id"]),
                file_name=str(book["file_name"]),
                sheet_name=str(book["sheet_name"]),
                data_start_row=int(book["data_start_row"]),
                footer_row=int(book["footer_row"]),
                sample_id_column=str(book["sample_id_column"]).upper(),
                age_column=str(book["age_column"]).upper(),
                era_labels=frozenset(
                    str(item).strip() for item in book.get("era_labels", [])
                ),
                fossil_classes=tuple(
                    FossilClassColumn(
                        column=str(item["column"]).upper(),
                        analyte_id=str(item["analyte_id"]),
                        measurement_channel=str(item["measurement_channel"]),
                        denominator_group=str(item["denominator_group"]),
                        used=bool(item["used"]),
                    )
                    for item in book["fossil_classes"]
                ),
            )
        )
    return Layout(
        workbooks=tuple(workbooks),
        standard_labels=frozenset(
            str(item).strip() for item in value["standard_labels"]
        ),
        compilation_labels=frozenset(
            str(item).strip() for item in value["compilation_labels"]
        ),
    )


def read_crosswalk(path: Path) -> list[dict[str, str]]:
    """Read the table that joins samples present in both workbooks."""
    with open(path, newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _json_value(value: Any) -> str:
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        raise ValueError("non-finite workbook value")
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _effective_value(formula_cell: Any, cached_cell: Any) -> Any:
    return cached_cell.value if formula_cell.data_type == "f" else formula_cell.value


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def value_type(value: Any, *, italic: bool) -> str:
    """Classify a cell value as blank, zero, positive (low signal if italic) or text."""
    if _is_blank(value):
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        if value == 0:
            return "reported_zero"
        if value > 0:
            return "reported_positive_low_signal" if italic else "reported_positive"
        return "reported_negative"
    return "qualified_text"


_RANGE_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*[-–—]\s*(\d+(?:\.\d+)?)\s*(?:Ma)?\s*$")
_UNCERTAINTY_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*±\s*(\d+(?:\.\d+)?)\s*(?:Ma)?\s*$")
_BOUND_RE = re.compile(r"^\s*([<>≤≥])\s*(\d+(?:\.\d+)?)\s*(?:Ma)?\s*$")
_APPROX_RE = re.compile(r"^\s*[~≈]\s*(\d+(?:\.\d+)?)\s*(?:Ma)?\s*$")
_EXACT_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*(?:Ma)?\s*$")


def parse_age(value: Any) -> AgeConstraint:
    """Parse an age cell into its kind and bounds, without substituting a midpoint."""
    raw = _text(value)
    if value is None or raw.strip() == "":
        return AgeConstraint(raw, "missing", None, None, None, None, "missing")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        return AgeConstraint(raw, "exact", number, number, number, "=", "parsed")
    match = _RANGE_RE.match(raw)
    if match:
        first, second = (float(item) for item in match.groups())
        return AgeConstraint(
            raw,
            "closed_interval",
            min(first, second),
            max(first, second),
            None,
            None,
            "parsed",
        )
    match = _UNCERTAINTY_RE.match(raw)
    if match:
        central, uncertainty = (float(item) for item in match.groups())
        return AgeConstraint(
            raw,
            "uncertainty_interval",
            central - uncertainty,
            central + uncertainty,
            central,
            "±",
            "parsed",
        )
    match = _BOUND_RE.match(raw)
    if match:
        operator, number_text = match.groups()
        number = float(number_text)
        if operator in {"<", "≤"}:
            return AgeConstraint(
                raw, "upper_bound", None, number, None, operator, "parsed"
            )
        return AgeConstraint(raw, "lower_bound", number, None, None, operator, "parsed")
    match = _APPROX_RE.match(raw)
    if match:
        return AgeConstraint(
            raw,
            "approximate",
            None,
            None,
            float(match.group(1)),
            raw.strip()[0],
            "approximate_no_interval",
        )
    match = _EXACT_RE.match(raw)
    if match:
        number = float(match.group(1))
        return AgeConstraint(raw, "exact", number, number, number, "=", "parsed")
    return AgeConstraint(raw, "unparsed", None, None, None, None, "unparsed")


def _row_values(formula_sheet: Any, cached_sheet: Any, row: int) -> list[Any]:
    return [
        _effective_value(
            formula_sheet.cell(row, column), cached_sheet.cell(row, column)
        )
        for column in range(1, formula_sheet.max_column + 1)
    ]


def _classify_rows(
    formula_sheet: Any, cached_sheet: Any, layout: WorkbookLayout
) -> dict[int, str]:
    rows: dict[int, str] = {}
    for row in range(layout.data_start_row, layout.footer_row):
        values = _row_values(formula_sheet, cached_sheet, row)
        sample_id = formula_sheet[f"{layout.sample_id_column}{row}"].value
        age = formula_sheet[f"{layout.age_column}{row}"].value
        if (
            _is_blank(sample_id)
            and _is_blank(age)
            and all(_is_blank(value) for value in values)
        ):
            kind = "blank"
        elif not _is_blank(sample_id) and not _is_blank(age):
            kind = "sample"
        elif not _is_blank(sample_id):
            label = str(sample_id).strip()
            kind = (
                "era_label"
                if label in layout.era_labels
                else "formation_or_basin_label"
            )
        else:
            kind = "aggregate_or_annotation"
        rows[row] = kind
    return rows


def _sample_context(
    formula_sheet: Any, layout: WorkbookLayout, row_kinds: Mapping[int, str]
) -> dict[int, tuple[str, str]]:
    era = ""
    formation = ""
    result: dict[int, tuple[str, str]] = {}
    for row, kind in row_kinds.items():
        label = _text(formula_sheet[f"{layout.sample_id_column}{row}"].value).strip()
        if kind == "era_label":
            era = label
            formation = ""
        elif kind == "formation_or_basin_label":
            formation = label
        elif kind == "sample":
            result[row] = (era, formation)
    return result


def _observations(
    formula_sheet: Any,
    cached_sheet: Any,
    layout: WorkbookLayout,
    row_kinds: Mapping[int, str],
) -> list[dict[str, Any]]:
    contexts = _sample_context(formula_sheet, layout, row_kinds)
    rows: list[dict[str, Any]] = []
    for row_number, kind in row_kinds.items():
        if kind != "sample":
            continue
        sample_id_cell = formula_sheet[f"{layout.sample_id_column}{row_number}"]
        age_cell = formula_sheet[f"{layout.age_column}{row_number}"]
        age = parse_age(age_cell.value)
        era, formation = contexts[row_number]
        sample_key = f"{layout.workbook_id}:{layout.sheet_name}:{row_number}"
        for fossil_class in layout.fossil_classes:
            formula_cell = formula_sheet[f"{fossil_class.column}{row_number}"]
            cached_cell = cached_sheet[f"{fossil_class.column}{row_number}"]
            if formula_cell.data_type == "f" and cached_cell.value is None:
                raise ValueError(
                    f"no stored value for formula cell {formula_cell.coordinate}"
                )
            value = _effective_value(formula_cell, cached_cell)
            italic = bool(formula_cell.font.italic)
            state = value_type(value, italic=italic)
            abundance = (
                float(value)
                if isinstance(value, (int, float)) and not isinstance(value, bool)
                else ""
            )
            rows.append(
                {
                    "workbook_id": layout.workbook_id,
                    "sheet_name": layout.sheet_name,
                    "sample_key": sample_key,
                    "source_row": row_number,
                    "source_sample_id_text": _text(sample_id_cell.value),
                    "source_sample_id_coordinate": sample_id_cell.coordinate,
                    "geological_era_text": era,
                    "formation_or_basin_text": formation,
                    "age_source_coordinate": age_cell.coordinate,
                    "age_raw_text": age.raw_text,
                    "age_kind": age.kind,
                    "age_lower_ma": "" if age.lower_ma is None else age.lower_ma,
                    "age_upper_ma": "" if age.upper_ma is None else age.upper_ma,
                    "age_central_ma": "" if age.central_ma is None else age.central_ma,
                    "age_operator": age.operator or "",
                    "age_parse_status": age.status,
                    "analyte_id": fossil_class.analyte_id,
                    "source_coordinate": formula_cell.coordinate,
                    "measurement_channel": fossil_class.measurement_channel,
                    "denominator_group": fossil_class.denominator_group,
                    "class_used": fossil_class.used,
                    "reported_abundance": abundance,
                    "value_type": state,
                    "qualifier_text": _text(value) if state == "qualified_text" else "",
                    "low_signal": italic and state.startswith("reported_positive"),
                    "cell_value": _json_value(cached_cell.value),
                }
            )
    return rows


def read_workbooks(layout: Layout, folder: Path) -> list[dict[str, Any]]:
    """Read every sample row of the two workbooks; formula cells give their stored values."""
    observations: list[dict[str, Any]] = []
    for book in layout.workbooks:
        path = Path(folder) / book.file_name
        formula_sheet = load_workbook(path, data_only=False, read_only=False)[
            book.sheet_name
        ]
        cached_sheet = load_workbook(path, data_only=True, read_only=False)[
            book.sheet_name
        ]
        row_kinds = _classify_rows(formula_sheet, cached_sheet, book)
        observations.extend(_observations(formula_sheet, cached_sheet, book, row_kinds))
    observations.sort(
        key=lambda row: (
            row["workbook_id"],
            row["source_row"],
            row["source_coordinate"],
        )
    )
    return observations


def normalize_sample_label(value: str) -> str:
    """Case-fold a sample label and keep only its letters and digits."""
    normalized = unicodedata.normalize("NFKC", value).casefold().strip()
    return "".join(character for character in normalized if character.isalnum())


def assign_samples(
    observations: Iterable[Mapping[str, Any]],
    layout: Layout,
    crosswalk: Iterable[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Mark standards and compiled rows, join samples present in both workbooks, and flag the values used."""
    source_rows = [dict(row) for row in observations]
    labels: dict[str, str] = {}
    for row in source_rows:
        labels.setdefault(
            str(row["sample_key"]), str(row["source_sample_id_text"]).strip()
        )

    joined: dict[str, tuple[str, tuple[str, str]]] = {}
    for row in crosswalk:
        if str(row.get("joined", "")).strip().lower() != "true":
            continue
        normalized = str(row["normalized_sample_label"]).strip()
        keys = (str(row["si1_sample_key"]).strip(), str(row["si2_sample_key"]).strip())
        for key in keys:
            if key not in labels or key in joined:
                raise ValueError(
                    f"crosswalk row does not match one workbook row: {key}"
                )
            if normalize_sample_label(labels[key]) != normalized:
                raise ValueError(
                    f"crosswalk label differs from the workbook label: {key}"
                )
            joined[key] = (normalized, keys)

    result: list[dict[str, Any]] = []
    for row in source_rows:
        sample_key = str(row["sample_key"])
        sample_label = str(row["source_sample_id_text"]).strip()
        if sample_label in layout.standard_labels:
            record_kind = "standard"
            included = False
            canonical_key = f"brocks_standard:{sample_key}"
            sample_record = "standard"
            reason = "standard_not_geological_sample"
        elif sample_label in layout.compilation_labels:
            record_kind = "compilation_context"
            included = False
            canonical_key = f"brocks_context:{sample_key}"
            sample_record = "compilation_row"
            reason = "compilation_context_only"
        else:
            record_kind = "named_physical_sample"
            included = True
            reason = "" if row["class_used"] else "source_interpretation_category"
            if sample_key in joined:
                normalized, _ = joined[sample_key]
                canonical_key = f"brocks_physical:{normalized}"
                sample_record = "both_workbooks"
            else:
                canonical_key = f"brocks_physical:source:{sample_key}"
                sample_record = "one_workbook"
        source_keys = joined[sample_key][1] if sample_key in joined else (sample_key,)
        class_used = bool(row["class_used"])
        value_used = included and class_used
        if included and not class_used:
            reason = "source_interpretation_category"
        row.update(
            {
                "canonical_sample_key": canonical_key,
                "record_kind": record_kind,
                "included_sample": included,
                "sample_record": sample_record,
                "source_rows": ";".join(sorted(source_keys)),
                "cross_workbook_join_applied": sample_key in joined,
                "class_used": class_used,
                "value_used": value_used,
                "exclusion_reason": "" if value_used else reason,
            }
        )
        result.append(row)
    result.sort(
        key=lambda row: (
            row["workbook_id"],
            row["source_row"],
            row["source_coordinate"],
        )
    )
    return result
